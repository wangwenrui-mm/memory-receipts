"""通用 Markdown → Word 生成器（复用工程报告那套排版）。

用法：
    python md2docx.py <输入.md> <输出.docx> [--title 标题] [--sub 副标题] [--info 行1|行2|...]
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from build_report_docx import (  # noqa: E402
    CJK_FONT, add_code_block, add_table, parse_md, set_run_font,
)
from docx.oxml.ns import qn  # noqa: E402


def add_image(doc, path: str, width_cm: float = 15.5):
    from docx.shared import Cm as _Cm
    p = Path(path)
    if not p.exists():
        return False
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(p), width=_Cm(width_cm))
    return True


def build(md_path: Path, out_path: Path, title: str, sub: str, info: list[str]):
    md = md_path.read_text(encoding="utf-8")
    doc = Document()

    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = CJK_FONT
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    set_run_font(r, CJK_FONT, 28, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(sub)
    set_run_font(r, CJK_FONT, 12.5, color=RGBColor(0x55, 0x55, 0x55))

    doc.add_paragraph()
    for line in info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, CJK_FONT, 10, color=RGBColor(0x66, 0x66, 0x66))

    doc.add_page_break()

    for kind, value in parse_md(md):
        if kind == "h":
            level, text = value
            h = doc.add_heading(level=min(level, 4))
            run = h.add_run(text.replace("**", "").replace("`", ""))
            size = {1: 20, 2: 15, 3: 12.5, 4: 11}.get(level, 11)
            set_run_font(run, CJK_FONT, size, bold=True,
                         color=RGBColor(0x1F, 0x3A, 0x5F) if level <= 2 else RGBColor(0x33, 0x33, 0x33))
        elif kind == "table":
            add_table(doc, value)
        elif kind == "code":
            add_code_block(doc, value)
        elif kind == "img":
            add_image(doc, value)
        elif kind == "li":
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(value.replace("**", "").replace("`", ""))
            set_run_font(r, CJK_FONT, 9.5)
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(value.replace("**", "").replace("`", ""))
            set_run_font(r, CJK_FONT, 9.5, color=RGBColor(0x55, 0x55, 0x55))
        else:
            text = value.replace("**", "").replace("`", "")
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            set_run_font(r, CJK_FONT, 10)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    kb = out_path.stat().st_size / 1024
    print(f"已生成: {out_path}  ({kb:.1f} KB)")


def main():
    args = sys.argv[1:]
    if len(args) < 2:
        print(__doc__)
        return 1
    md_path, out_path = Path(args[0]), Path(args[1])
    title = "文档"
    sub = ""
    info: list[str] = []
    if "--title" in args:
        title = args[args.index("--title") + 1]
    if "--sub" in args:
        sub = args[args.index("--sub") + 1]
    if "--info" in args:
        info = args[args.index("--info") + 1].split("|")
    build(md_path, out_path, title, sub, info)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
